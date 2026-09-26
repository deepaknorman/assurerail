#!/usr/bin/env python3
"""Generate fictional NBFC inputs, never assessment decisions or released offers."""
import argparse
import csv
import json
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AS_OF = date(2026, 9, 15)
HEADERS = ["loan_id", "party_id", "party_role", "principal_minor"]
BOOKS = [
    ("A", "demo-nbfc-ev-001", "AssureRail Synthetic EV Finance NBFC", "SYN-EV-2W-2025", "Electric two-wheeler", 24, 6, 3, "duplicate"),
    ("B", "demo-nbfc-ev-001", "AssureRail Synthetic EV Finance NBFC", "SYN-EV-3W-2025", "Electric three-wheeler", 18, 4, 2, "balance_and_role"),
    ("C", "demo-nbfc-ev-002", "Synthetic Summit Vehicle Finance NBFC", "SYN-EV-2W-2025", "Electric two-wheeler", 20, 5, 3, "clean_tape"),
]

LARGE_BOOKS = [
    ("A", "demo-nbfc-ev-001", "AssureRail Synthetic EV Finance NBFC", "SYN-EV-2W-2025", "Electric two-wheeler", 1200, 300, 120, "duplicate"),
    ("B", "demo-nbfc-ev-001", "AssureRail Synthetic EV Finance NBFC", "SYN-EV-3W-2025", "Electric three-wheeler", 800, 200, 80, "balance_and_role"),
    ("C", "demo-nbfc-ev-002", "Synthetic Summit Vehicle Finance NBFC", "SYN-EV-2W-2025", "Electric two-wheeler", 600, 150, 60, "clean_tape"),
    ("D", "demo-nbfc-ev-002", "Synthetic Summit Vehicle Finance NBFC", "SYN-EV-3W-2025", "Electric three-wheeler", 400, 100, 40, "clean_tape"),
]
LARGE_TARGETS = {"A": 20_000_000_000, "B": 15_000_000_000, "C": 12_000_000_000, "D": 8_000_000_000}  # INR paise


def allocate_exact(weights, target):
    """Largest-remainder allocation preserves varied ticket sizes and exact integer totals."""
    total = sum(weights)
    allocations = [weight * target // total for weight in weights]
    remainders = [weight * target % total for weight in weights]
    for index in sorted(range(len(weights)), key=lambda i: (-remainders[i], i))[:target - sum(allocations)]:
        allocations[index] += 1
    assert sum(allocations) == target and min(allocations) > 0
    return allocations


def rounded(value):
    return int(Decimal(value).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def csv_file(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        writer.writerows(rows)


def create(output, profile="small"):
    books = LARGE_BOOKS if profile == "55cr" else BOOKS
    workbook_name = "AssureRail_Synthetic_NBFC_55Cr_Step1.xlsx" if profile == "55cr" else "AssureRail_Synthetic_NBFC_Step1.xlsx"
    output.mkdir(parents=True, exist_ok=True)
    masters, links, scenarios, book_summaries, tapes, manifest = [], [], [], [], [], []
    master_headers = ["classification", "seller_institution_id", "seller_name", "loanbook_ref", "loan_id", "borrower_id", "borrower_display_name", "vehicle_type", "state", "origination_date", "as_of_date", "original_principal_inr", "principal_outstanding_inr", "annual_interest_rate_percent", "original_term_months", "instalments_paid", "scheduled_instalment_inr", "days_past_due", "missed_instalments"]
    for book_index, (tag, seller, name, book_ref, vehicle, loan_count, co_count, linked_count, scenario) in enumerate(books):
        clean, balances, model = [], [], []
        for index in range(1, loan_count + 1):
            rate, term = Decimal(14 + index % 5) / 1200, 36 + (index % 2) * 12
            originated = date(2025, 1 + index % 9, 15)
            elapsed = (AS_OF.year - originated.year) * 12 + AS_OF.month - originated.month
            missed = 0 if index % 6 else 1 + ((index // 6) % 3)
            paid = elapsed - missed
            factor = (1 + rate) ** paid - rate / (1 - (1 + rate) ** -term) * ((1 + rate) ** paid - 1) / rate
            base_minor = (160_000 + book_index * 70_000 + (index % 9) * 17_500) * 100
            model.append((factor, rounded(Decimal(base_minor) * factor)))
        allocated = allocate_exact([weight for _, weight in model], LARGE_TARGETS[tag]) if profile == "55cr" else None
        for index in range(1, loan_count + 1):
            loan_id, borrower_id = f"SYN-{tag}-LN-{index:04d}", f"SYN-{tag}-BR-{index:04d}"
            original_minor = rounded(Decimal(allocated[index - 1]) / model[index - 1][0]) if allocated else (160_000 + book_index * 70_000 + (index % 9) * 17_500) * 100
            rate_percent, term = Decimal(14 + index % 5), 36 + (index % 2) * 12
            monthly_rate = rate_percent / 1200
            emi_minor = Decimal(original_minor) * monthly_rate / (1 - (1 + monthly_rate) ** -term)
            originated = date(2025, 1 + (index % 9), 15)
            elapsed = (AS_OF.year - originated.year) * 12 + AS_OF.month - originated.month
            missed = 0 if index % 6 else 1 + ((index // 6) % 3)
            paid = elapsed - missed
            outstanding = rounded(Decimal(original_minor) * (1 + monthly_rate) ** paid - emi_minor * ((1 + monthly_rate) ** paid - 1) / monthly_rate)
            if allocated:
                assert outstanding == allocated[index - 1], "amortisation must reproduce the allocated balance"
            balances.append(outstanding)
            clean.append([loan_id, borrower_id, "BORROWER", str(outstanding)])
            masters.append(["SYNTHETIC_ONLY", seller, name, book_ref, loan_id, borrower_id, f"Synthetic Borrower {tag}{index:04d}", vehicle, ["Maharashtra", "Karnataka", "Tamil Nadu"][index % 3], originated.isoformat(), AS_OF.isoformat(), original_minor / 100, outstanding / 100, float(rate_percent), term, paid, rounded(emi_minor) / 100, missed * 30 + (index % 9) if missed else 0, missed])
            if index <= co_count:
                clean.append([loan_id, f"SYN-{tag}-CO-{index:04d}", "CO_BORROWER", str(outstanding)])
            if index <= linked_count:
                clean.append([loan_id, f"SYN-{tag}-GU-{index:04d}", "LINKED_PARTY", str(outstanding)])
        initial = [row.copy() for row in clean]
        defects = []
        if scenario == "duplicate":
            borrowers = [row for row in clean if row[2] == "BORROWER"]
            for row in borrowers[:max(1, loan_count // 100)]:
                initial.append(row.copy())
                defects.append({"csvRow": len(initial) + 1, "code": "DUPLICATE_PAIR", "loanId": row[0], "description": "Borrower pair appears twice; principal must not be counted twice."})
        elif scenario == "balance_and_role":
            for position in [i for i, row in enumerate(initial) if row[2] == "CO_BORROWER"][:max(1, loan_count // 100)]:
                initial[position][3] = str(int(initial[position][3]) + 10000)
                defects.append({"csvRow": position + 2, "code": "INCONSISTENT_LOAN_BALANCE", "loanId": initial[position][0], "description": "Co-borrower repeats this loan's outstanding balance with a deliberate INR 100 mismatch."})
            for position in [i for i, row in enumerate(initial) if row[2] == "LINKED_PARTY"][:max(1, loan_count // 100)]:
                initial[position][2] = "GUARANTOR"
                defects.append({"csvRow": position + 2, "code": "INVALID_RECORD", "loanId": initial[position][0], "description": "Unmapped source role GUARANTOR; canonical input requires LINKED_PARTY."})
        prefix = f"book-{tag.lower()}"
        current_file = f"upload-step1/{prefix}-initial.csv"
        reference_file = f"reference-for-step2/{prefix}-corrected.csv"
        csv_file(output / current_file, HEADERS, initial)
        csv_file(output / reference_file, HEADERS, clean)
        links.extend([["SYNTHETIC_ONLY", seller, book_ref, *row] for row in clean])
        proposed = sum(balances)  # Synthetic declaration at par, not a valuation or buyer offer.
        manifest.append({"key": [seller, book_ref], "label": f"Book {tag}", "sellerInstitutionId": seller, "sellerName": name, "bookRef": book_ref, "assetFamily": "VEHICLE_EV", "asOfDate": AS_OF.isoformat(), "loanCount": loan_count, "primaryPairCount": loan_count + co_count, "linkedPartyCount": linked_count, "sellerProposedConsiderationMinor": str(proposed), "aggregateProgrammeConsiderationMinor": str(sum(LARGE_TARGETS.values()) if profile == "55cr" else proposed), "considerationBasis": "SYNTHETIC_PAR_DECLARATION_NOT_VALUATION", "principalMinor": str(proposed), "initialFile": current_file, "referenceFile": reference_file, "initialExpectedStatus": "RECORD_EXCEPTIONS" if defects else "MATCHED", "referenceExpectedStatus": "MATCHED", "deliberateDefects": defects, "hostedInstitutionAvailableAtPreparation": seller == "demo-nbfc-ev-001"})
        book_summaries.append([tag, seller, name, book_ref, loan_count, loan_count + co_count, linked_count, proposed / 100, "RECORD_EXCEPTIONS" if defects else "MATCHED", "Supporting evidence missing", current_file])
        tapes.append((f"Book {tag} upload", initial))
        for defect in defects:
            scenarios.append([tag, current_file, defect["csvRow"], defect["code"], defect["loanId"], defect["description"]])
    csv_file(output / "reference-loan-master.csv", master_headers, masters)
    csv_file(output / "reference-party-links.csv", ["classification", "seller_institution_id", "loanbook_ref", *HEADERS], links)
    seller_totals = {}
    for book in manifest:
        seller_totals[book["sellerInstitutionId"]] = seller_totals.get(book["sellerInstitutionId"], 0) + int(book["principalMinor"])
    payload = {"profile": profile, "workbookFile": workbook_name, "portfolioPrincipalMinor": str(sum(seller_totals.values())), "sellerPrincipalMinor": {key: str(value) for key, value in seller_totals.items()}, "schemaVersion": 1, "classification": "SYNTHETIC_ONLY", "purpose": "STEP_1_AUTOMATIC_READINESS_INPUTS", "asOfDate": AS_OF.isoformat(), "books": manifest, "missingEvidenceFamilies": ["LOAN_AGREEMENT", "SECURITY_DOCUMENT", "REPAYMENT_HISTORY", "KYC_AUTHORITY", "INSURANCE_COLLATERAL"], "assessmentResultsIncluded": False, "finalOffersIncluded": False}
    (output / "input-manifest.json").write_text(json.dumps(payload, indent=2) + "\n")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "AssureRail synthetic test-input generator"
    wb.properties.created = wb.properties.modified = datetime(2026, 9, 26)
    sections = [
        ("Start here", ["Item", "Instruction"], [
            ["Classification", "SYNTHETIC ONLY. Fictional institutions, people and balances; no customer data."],
            ["Step 1", "Upload one canonical CSV from upload-step1 as LOAN_TAPE in its own seller/book engagement."],
            ["Excel status", "Reference workbook. Automatic loan-tape metrics currently require CSV; XLSX needs mapping."],
            ["Currency", "principal_minor is INR paise. INR 1 = 100 paise. Balances repeat across parties; count each loan once."],
            ["Scope counts", "Primary units = BORROWER + CO_BORROWER pairs. LINKED_PARTY pairs are separate."],
            ["Starting files", "A: duplicate pair. B: inconsistent balance and unmapped party role. C/D where present: clean tapes."],
            ["Missing evidence", "All books lack the five supporting document families; a clean tape is not a ready portfolio."],
            ["Step 2 reference", "Corrected tapes are supplied separately for later reassessment. Do not upload initial and corrected together."],
            ["Seller isolation", "Books C/D use a second institution not yet provisioned on the demo host. A and C intentionally reuse bookRef."],
            ["Cohort identity", "Use (sellerInstitutionId, bookRef), never bookRef alone. This is presentation grouping, not billing allocation."],
            ["Money", "Declared consideration at synthetic par is a test input, not valuation, approval or final offer. Large profile: INR 55 crore programme, 35 + 20 crore sellers."],
            ["AI scale boundary", "Full tapes are checked deterministically. The current worker AI budget is exceeded at 55 crore fixture volumes; complete AI review is not claimed." if profile == "55cr" else "Supporting-document completion and actual provider processing remain required."],
            ["Loan model", "Monthly amortisation; missed instalments delay scheduled principal reduction. Excludes fees and penalty interest."],
        ]),
        ("Book summary", ["Book", "Seller institution", "Seller", "Loanbook", "Loans", "Primary pairs", "Linked pairs", "Outstanding INR", "Expected tape check", "Overall readiness caveat", "Upload CSV"], book_summaries),
        ("Loan master", master_headers, masters),
        *([("Party links", ["classification", "seller_institution_id", "loanbook_ref", *HEADERS], links)] if profile == "small" else []),
        *[(title, HEADERS, rows) for title, rows in tapes],
        ("Deliberate defects", ["Book", "File", "CSV row", "Code", "Loan", "Reason"], scenarios),
    ]
    for title, headers, rows in sections:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="143B50")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EEF5F8")
        for index, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(index)].width = min(42, max(18, len(header) + 2))
            if "inr" in header.lower():
                for cell in list(ws.columns)[index - 1][1:]:
                    cell.number_format = '#,##0.00'
        if title == "Start here":
            ws.column_dimensions["B"].width = 110
            for row in range(2, ws.max_row + 1):
                ws.row_dimensions[row].height = 34
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.print_title_rows = "1:1"
    wb.save(output / workbook_name)
    print(json.dumps({"output": str(output), "books": len(manifest), "loans": len(masters), "classification": "SYNTHETIC_ONLY"}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--profile", choices=["small", "55cr"], default="small")
    args = parser.parse_args()
    create(args.output_dir, args.profile)
