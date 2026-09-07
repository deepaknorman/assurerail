#!/usr/bin/env python3
"""Build the governed AssureRail pilot pricing workbook from approved planning assumptions."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/gtm/templates/AssureRail_Pilot_Pricing_Worksheet.xlsx"

NAVY = "0E2836"
GREEN = "157A63"
MINT = "DDF3EC"
PALE = "F4F7F8"
AMBER = "F7E8BD"
RED = "F7D7D7"
WHITE = "FFFFFF"
GREY = "5D6B73"
THIN = Side(style="thin", color="CFD8DC")


def title(ws, text, subtitle):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = text
    ws["A1"].font = Font(name="Arial", size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:F2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color=GREY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 32


def section(ws, row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    cell = ws.cell(row, 1, text)
    cell.font = Font(name="Arial", size=11, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=GREEN)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 23


def style_table(ws, start_row, end_row, start_col=1, end_col=6):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=NAVY)
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_quote_sheet(wb):
    ws = wb.active
    ws.title = "Quote Builder"
    title(ws, "AssureRail pilot pricing worksheet", "CONFIDENTIAL DRAFT · Version 1.0 · 7 September 2026 · not a quote until approved")
    section(ws, 4, "Engagement and approval")
    rows = [
        ("Customer legal name", "", "Input", "Required"),
        ("AssureRail contracting entity", "", "Input", "Must exist and be authorised before signature"),
        ("Engagement mode", "REPLAY", "Input", "DEMONSTRATION / REPLAY / SHADOW / CONTROLLED_LIVE"),
        ("Route", "PTC", "Input", "DA / PTC / BOTH"),
        ("Representation", "CONVENTIONAL", "Input", "CONVENTIONAL / TOKENISED"),
        ("Commercial owner", "", "Input", "Named approver"),
        ("Approval date", "", "Input", "Required before external use"),
        ("Quote valid until", "", "Input", "Required before external use"),
    ]
    for index, (label, value, kind, note) in enumerate(rows, 5):
        ws.cell(index, 1, label)
        ws.cell(index, 2, value)
        ws.cell(index, 3, kind)
        ws.cell(index, 4, note)
        ws.merge_cells(start_row=index, start_column=4, end_row=index, end_column=6)
        ws.cell(index, 2).fill = PatternFill("solid", fgColor=MINT)
    style_table(ws, 5, 12)

    mode = DataValidation(type="list", formula1='"DEMONSTRATION,REPLAY,SHADOW,CONTROLLED_LIVE"')
    route = DataValidation(type="list", formula1='"DA,PTC,BOTH"')
    representation = DataValidation(type="list", formula1='"CONVENTIONAL,TOKENISED"')
    for validation, cell in [(mode, "B7"), (route, "B8"), (representation, "B9")]:
        ws.add_data_validation(validation)
        validation.add(ws[cell])

    section(ws, 14, "Fee inputs and exact calculation")
    fee_rows = [
        ("Defined transaction fee base (INR)", 0, "Input", "Define inclusions and exclusions in words in the order form"),
        ("Platform rate (bps)", '=IF(AND(B7="CONTROLLED_LIVE",B8="PTC",B9="CONVENTIONAL"),30,IF(AND(B7="CONTROLLED_LIVE",B8="PTC",B9="TOKENISED"),50,0))', "Formula / override", "Replay/shadow default to zero variable fee; 30/50 bps applies only to a separately approved controlled-live PTC quote"),
        ("Fixed discovery/replay fee (INR)", 0, "Input", "Prospect-specific; zero does not imply waiver"),
        ("Implementation/integration fee (INR)", 0, "Input", "Scope, milestones and change control required"),
        ("Lifecycle/monitoring fee (INR)", 0, "Input", "State period, coverage and service level"),
        ("Third-party charges (INR)", 0, "Input", "State provider and pass-through/markup treatment"),
        ("Minimum platform fee (INR)", 0, "Input", "Zero means no minimum in this calculation"),
        ("Maximum platform fee (INR)", 0, "Input", "Zero means no cap in this calculation"),
        ("Calculated variable fee (INR)", "=ROUND(B15*B16/10000,2)", "Formula", "Exact bps calculation"),
        ("Platform fee after minimum/cap (INR)", "=ROUND(MIN(IF(B22>0,B22,1E+30),MAX(B21,B23)),2)", "Formula", "Minimum and cap apply only if entered"),
        ("Subtotal before tax (INR)", "=ROUND(SUM(B17:B20)+B24,2)", "Formula", "Excludes GST and other statutory levies"),
        ("GST rate", 0.18, "Input", "Confirm applicable tax treatment"),
        ("GST amount (INR)", "=ROUND(B25*B26,2)", "Formula", "Illustrative pending tax review"),
        ("Total invoice illustration (INR)", "=ROUND(B25+B27,2)", "Formula", "Not an invoice or revenue-recognition conclusion"),
    ]
    for index, (label, value, kind, note) in enumerate(fee_rows, 15):
        ws.cell(index, 1, label)
        ws.cell(index, 2, value)
        ws.cell(index, 3, kind)
        ws.cell(index, 4, note)
        ws.merge_cells(start_row=index, start_column=4, end_row=index, end_column=6)
        ws.cell(index, 2).fill = PatternFill("solid", fgColor=MINT if kind.startswith("Input") or "override" in kind else PALE)
        ws.cell(index, 2).number_format = "₹#,##0.00"
    ws["B16"].number_format = '0 "bps"'
    ws["B26"].number_format = "0.00%"
    style_table(ws, 15, 28)
    ws.conditional_formatting.add("B16", CellIsRule(operator="greaterThan", formula=["50"], fill=PatternFill("solid", fgColor=RED)))

    section(ws, 30, "Commercial terms — must be completed")
    terms = [
        "Fee-base definition", "Earned/billing trigger", "Minimum and cap operation",
        "Cancellation/failure/reversal/refund treatment", "Invoice timing and payment days",
        "Third-party charges and taxes", "Discount approval and expiry", "Change-control rate",
        "Assurance-provider fee stated separately", "Contract and SOW version reference",
    ]
    for index, label in enumerate(terms, 31):
        ws.cell(index, 1, label)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=5)
        ws.cell(index, 2, "")
        ws.cell(index, 2).fill = PatternFill("solid", fgColor=MINT)
        ws.cell(index, 6, "OPEN")
        ws.cell(index, 6).fill = PatternFill("solid", fgColor=AMBER)
    style_table(ws, 31, 40)

    for column, width in {"A": 34, "B": 23, "C": 18, "D": 24, "E": 22, "F": 18}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A30:F40"
    ws.print_area = "A1:F40"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 2
    return ws


def build_rate_card(wb):
    ws = wb.create_sheet("Rate Governance")
    title(ws, "Rate governance", "Planning assumptions and quotation boundaries — not public list prices")
    headers = ["Service", "Mode", "Route", "Representation", "Planning basis", "Boundary"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=GREEN)
    data = [
        ("Discovery workshop", "Any", "Any", "Any", "Fixed / approval", "May be waived or credited only in writing"),
        ("Completed-deal replay", "REPLAY", "DA or PTC", "CONVENTIONAL", "Fixed scoped fee", "Set after evidence-family and complexity review"),
        ("Shadow evaluation", "SHADOW", "DA or PTC", "CONVENTIONAL", "Fixed programme + integration", "No execution implication"),
        ("Conventional DA", "Future controlled-live", "DA", "CONVENTIONAL", "TBD", "Do not infer PTC rate"),
        ("Conventional PTC", "Future approved service", "PTC", "CONVENTIONAL", "30 bps", "Fee base, trigger, minimum/cap and refunds must be defined"),
        ("Tokenised PTC", "Future separately gated", "PTC", "TOKENISED", "50 bps", "Not currently available; route/legal/operating gates apply"),
        ("Lifecycle/monitoring", "By scope", "DA or PTC", "Any approved", "Annual/programme/case tier", "Define coverage, frequency and providers"),
        ("Independent assurance", "Separate appointment", "Any", "Any", "Separate provider fee", "Never silently bundled or outcome-dependent"),
    ]
    for row, values in enumerate(data, 5):
        for col, value in enumerate(values, 1):
            ws.cell(row, col, value)
    style_table(ws, 5, 12)
    for column, width in {"A": 27, "B": 24, "C": 18, "D": 20, "E": 24, "F": 48}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:F12"
    return ws


def build_approval_sheet(wb):
    ws = wb.create_sheet("Approval Gate")
    title(ws, "Proposal approval gate", "Every mandatory row must be CLOSED before external use")
    headers = ["Gate", "Owner", "Status", "Evidence/reference", "Checked at", "Notes"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=GREEN)
    gates = [
        ("Contracting entity and signatory", "Legal/commercial"),
        ("Engagement mode and function boundary", "Product/legal"),
        ("Fee base and trigger", "Commercial"),
        ("Minimum/cap/refund treatment", "Commercial/finance"),
        ("Data purpose and authority", "Privacy/customer owner"),
        ("Security and environment scope", "Security/technical"),
        ("External provider charges and roles", "Operations/commercial"),
        ("Claims and availability wording", "Product/founder"),
        ("Tax review", "Finance/tax"),
        ("Agreement/SOW versions matched", "Legal/operations"),
    ]
    validation = DataValidation(type="list", formula1='"OPEN,CLOSED,NOT_APPLICABLE"')
    ws.add_data_validation(validation)
    for row, (gate, owner) in enumerate(gates, 5):
        ws.cell(row, 1, gate)
        ws.cell(row, 2, owner)
        ws.cell(row, 3, "OPEN")
        validation.add(ws.cell(row, 3))
        ws.cell(row, 3).fill = PatternFill("solid", fgColor=AMBER)
    style_table(ws, 5, 14)
    ws.conditional_formatting.add("C5:C14", CellIsRule(operator="equal", formula=['"CLOSED"'], fill=PatternFill("solid", fgColor=MINT)))
    for column, width in {"A": 34, "B": 24, "C": 20, "D": 34, "E": 20, "F": 36}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    return ws


def main():
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    build_quote_sheet(wb)
    build_rate_card(wb)
    build_approval_sheet(wb)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.protection = Protection(locked=False)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.title = "AssureRail Pilot Pricing Worksheet"
    wb.properties.subject = "Confidential draft pricing and approval control"
    wb.properties.creator = "AssureRail"
    wb.properties.description = "Planning workbook; not a quote until the approval gate is closed."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
